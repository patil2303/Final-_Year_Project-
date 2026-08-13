import io
import re
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import List, Dict, Any, Optional

def convert_sections_to_dataframes(sections: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Converts section matrix data into Pandas DataFrames with cleaned numeric types."""
    result = []

    for sec in sections:
        headers = sec.get("headers", [])
        raw_rows = sec.get("rows", [])

        # Build clean data matrix
        df_matrix = []
        for r in raw_rows:
            row_vals = []
            for cell in r:
                val = cell.get("value", "")
                if val == "":
                    row_vals.append(None)
                else:
                    row_vals.append(val)
            df_matrix.append(row_vals)

        # Create DataFrame
        if not headers and df_matrix:
            headers = [f"Col_{i+1}" for i in range(len(df_matrix[0]))]

        df = pd.DataFrame(df_matrix, columns=headers if headers else None)
        
        # Convert numeric columns where possible
        for col in df.columns:
            df[col] = pd.to_numeric(df[col], errors='ignore')

        result.append({
            "title": sec.get("title", "Section"),
            "dataframe": df
        })

    return result

def attach_metadata_to_sections(sections: List[Dict[str, Any]], metadata: Dict[str, Any]) -> List[Dict[str, Any]]:
    """
    Prepends student metadata columns (Student Name, PRN Number, Roll No, Branch, Division, Semester)
    to section table headers and rows so that student details and marks exist in a single row.
    """
    if not metadata or not any(str(v).strip() for v in metadata.values()):
        return sections

    meta_cols = [
        ("Student Name", metadata.get("student_name", "")),
        ("PRN Number", metadata.get("prn", "")),
        ("Roll No", metadata.get("roll_no", "")),
        ("Branch", metadata.get("branch", "")),
        ("Division", metadata.get("division", "")),
        ("Semester", metadata.get("semester", "")),
        ("Subject", metadata.get("subject", ""))
    ]

    active_meta_cols = [(label, str(val).strip()) for label, val in meta_cols if str(val).strip()]
    if not active_meta_cols:
        return sections

    meta_headers = [label for label, _ in active_meta_cols]

    updated_sections = []
    for sec in sections:
        sec_copy = dict(sec)
        orig_headers = list(sec.get("headers", []))
        orig_rows = sec.get("rows", [])

        # Check if metadata columns are already present
        meta_indices = {}
        for idx, h in enumerate(orig_headers):
            h_norm = str(h).strip().lower()
            for label, _ in active_meta_cols:
                if h_norm == label.lower():
                    meta_indices[label] = idx

        if meta_indices:
            # Metadata columns already exist; update their row values with latest user edits!
            new_rows = []
            for r_idx, r_cells in enumerate(orig_rows):
                new_r_cells = list(r_cells)
                first_cell_text = str(r_cells[0].get("value", "")).lower() if r_cells else ""
                is_max_marks_row = "max" in first_cell_text or "q.no" in first_cell_text

                for label, val in active_meta_cols:
                    if label in meta_indices:
                        col_i = meta_indices[label]
                        cell_val = "Max Marks" if is_max_marks_row and label == "Student Name" else (val if not is_max_marks_row else "")
                        if col_i < len(new_r_cells):
                            new_r_cells[col_i] = {
                                "text": str(cell_val),
                                "value": str(cell_val),
                                "is_number": False,
                                "bbox": None,
                                "confidence": 1.0,
                                "classifier": "Student Metadata Header"
                            }
                new_rows.append(new_r_cells)

            sec_copy["rows"] = new_rows
            sec_copy["metadata"] = metadata
            updated_sections.append(sec_copy)
            continue

        new_headers = meta_headers + orig_headers

        new_rows = []
        for r_idx, r_cells in enumerate(orig_rows):
            new_r_cells = []
            
            # Check if this row is a Max Marks header row vs Marks Awarded row
            first_cell_text = str(r_cells[0].get("value", "")).lower() if r_cells else ""
            is_max_marks_row = "max" in first_cell_text or "q.no" in first_cell_text

            for label, val in active_meta_cols:
                cell_val = "Max Marks" if is_max_marks_row and label == "Student Name" else (val if not is_max_marks_row else "")
                new_r_cells.append({
                    "text": str(cell_val),
                    "value": str(cell_val),
                    "is_number": False,
                    "bbox": None,
                    "confidence": 1.0,
                    "classifier": "Student Metadata Header"
                })
            
            new_r_cells.extend(r_cells)
            new_rows.append(new_r_cells)

        sec_copy["headers"] = new_headers
        sec_copy["rows"] = new_rows
        sec_copy["columns_count"] = len(new_headers)
        sec_copy["metadata"] = metadata
        updated_sections.append(sec_copy)

    return updated_sections

def export_to_excel_bytes(sections: List[Dict[str, Any]], metadata: Optional[Dict[str, Any]] = None) -> bytes:
    """
    Generates a beautifully styled Excel (.xlsx) file in bytes.
    Includes custom header fills, cell number formatting, auto-column sizing,
    and prepended student metadata columns for single-row Excel export.
    """
    meta_to_use = metadata
    if not meta_to_use:
        for s in sections:
            if s.get("metadata"):
                meta_to_use = s.get("metadata")
                break

    if meta_to_use:
        sections = attach_metadata_to_sections(sections, meta_to_use)

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    section_fill = PatternFill(start_color="475569", end_color="475569", fill_type="solid")
    section_font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")

    grid_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # 1. Master Sheet (All Sections Stacked)
    ws_master = wb.create_sheet(title="Summary")
    current_row = 1

    for sec in sections:
        title = sec.get("title", "Section")
        headers = sec.get("headers", [])
        rows = sec.get("rows", [])

        # Section Header Banner
        num_cols = max(len(headers), 1)
        ws_master.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=num_cols)
        banner_cell = ws_master.cell(row=current_row, column=1, value=title.upper())
        banner_cell.fill = section_fill
        banner_cell.font = section_font
        banner_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws_master.row_dimensions[current_row].height = 28
        current_row += 1

        # Table Column Headers
        if headers:
            for c_idx, h_text in enumerate(headers, 1):
                cell = ws_master.cell(row=current_row, column=c_idx, value=str(h_text))
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = grid_border
            ws_master.row_dimensions[current_row].height = 22
            current_row += 1

        # Table Rows
        for r_idx, r_cells in enumerate(rows, 1):
            for c_idx, c_info in enumerate(r_cells, 1):
                val = c_info.get("value", "")
                cell = ws_master.cell(row=current_row, column=c_idx)

                # Set numeric or text value
                if isinstance(val, (int, float)):
                    cell.value = val
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if isinstance(val, float):
                        cell.number_format = '#,##0.00'
                    else:
                        cell.number_format = '#,##0'
                else:
                    cell.value = str(val) if val is not None else ""
                    cell.alignment = Alignment(horizontal="left", vertical="center")

                cell.border = grid_border
            ws_master.row_dimensions[current_row].height = 20
            current_row += 1

        # Gap between sections
        current_row += 2

    # Auto-adjust column widths for Summary sheet
    for col in ws_master.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            max_len = max(max_len, len(val_str))
        ws_master.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # 2. Separate Sheets for each section if more than 1 section
    if len(sections) > 1:
        for idx, sec in enumerate(sections, 1):
            raw_title = sec.get("title", f"Section_{idx}")
            sheet_title = re.sub(r'[\:\\/\?\*\[\]]', '_', raw_title)[:30] or f"Sheet_{idx}"
            
            ws = wb.create_sheet(title=sheet_title)
            headers = sec.get("headers", [])
            rows = sec.get("rows", [])

            # Write headers
            if headers:
                for c_idx, h_text in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=c_idx, value=str(h_text))
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = grid_border

            # Write rows
            start_r = 2 if headers else 1
            for r_idx, r_cells in enumerate(rows, start_r):
                for c_idx, c_info in enumerate(r_cells, 1):
                    val = c_info.get("value", "")
                    cell = ws.cell(row=r_idx, column=c_idx)
                    if isinstance(val, (int, float)):
                        cell.value = val
                        cell.alignment = Alignment(horizontal="right")
                    else:
                        cell.value = str(val) if val is not None else ""
                    cell.border = grid_border

            # Auto-adjust column widths (guard against empty sheets)
            for col in ws.columns:
                col_cells = list(col)
                if not col_cells:
                    continue
                max_len = max((len(str(cell.value or '')) for cell in col_cells), default=0)
                col_letter = get_column_letter(col_cells[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

def _sanitize_csv_value(val) -> str:
    """Sanitizes a value to prevent CSV formula injection in spreadsheet applications."""
    s = str(val) if val is not None else ""
    # If the value starts with a formula trigger character, prefix with a single quote
    if s and s[0] in ('=', '+', '-', '@', '\t', '\r', '\n'):
        # Check if it's a legitimate negative number
        try:
            float(s)
            return s  # legitimate number like "-5.3"
        except (ValueError, TypeError):
            return "'" + s  # escape formula
    return s

def export_to_csv_string(sections: List[Dict[str, Any]], metadata: Optional[Dict[str, Any]] = None) -> str:
    """Generates clean CSV text with section boundary titles and prepended student metadata."""
    meta_to_use = metadata
    if not meta_to_use:
        for s in sections:
            if s.get("metadata"):
                meta_to_use = s.get("metadata")
                break

    if meta_to_use:
        sections = attach_metadata_to_sections(sections, meta_to_use)

    lines = []
    for sec in sections:
        title = sec.get("title", "Section")
        lines.append(f"# --- SECTION: {title} ---")

        headers = sec.get("headers", [])
        if headers:
            lines.append(",".join(f'"{ _sanitize_csv_value(h)}"' for h in headers))

        rows = sec.get("rows", [])
        for r in rows:
            row_vals = []
            for cell in r:
                val = cell.get("value", "")
                sanitized = _sanitize_csv_value(val)
                row_vals.append(f'"{sanitized}"' if isinstance(val, str) else str(sanitized))
            lines.append(",".join(row_vals))

        lines.append("")  # Blank line separator

    return "\n".join(lines)
